#!/usr/bin/env python3
"""Build the master workbook. Refuse to overwrite later manual additions."""
import hashlib
import json
import shutil
import time
import urllib.request
import zipfile
from collections import Counter
from concurrent.futures import ThreadPoolExecutor, as_completed
from pathlib import Path
from PIL import Image
import xlsxwriter
from export_local_controller_excel import GENRE_VI, qualifies, unique_games
from export_excel import clean, nested, profile_row, HEADERS as DETAIL_HEADERS

ROOT = Path('/workspace')
DATA = ROOT / 'game-steam/data/game-profiles.json'
OUTPUT = ROOT / 'Thu-vien-game-TONG.xlsx'
CACHE = ROOT / 'game-steam/exports/master-covers'
REPORT = ROOT / 'game-steam/exports/master-image-report.json'
HEADERS = ['STT', 'Ảnh', 'Tên game', 'Thể loại', 'Tay cầm', 'Local Co-op', 'Chia sẻ-Chia màn hình', 'Online Co-op', 'Windows tối thiểu', 'Nền tảng', 'Link Steam', 'URL ảnh', 'Mã hồ sơ', 'Trạng thái xác minh']

def path_for(url):
    return CACHE / (hashlib.sha256(url.encode()).hexdigest() + '.jpg')

def fetch(url):
    path = path_for(url)
    if path.exists():
        return str(path), url
    for attempt in range(3):
        try:
            with urllib.request.urlopen(urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0'}), timeout=25) as response:
                with Image.open(response) as image:
                    image = image.convert('RGB')
                    image.thumbnail((160, 160))
                    image.save(path, 'JPEG', quality=85)
            return str(path), url
        except Exception:
            if attempt == 2:
                raise
            time.sleep(attempt + 1)

def flag(value):
    return 'Có' if value is True else 'Không' if value is False else 'Chưa xác định'

def main():
    if OUTPUT.exists():
        raise RuntimeError('File tổng đã tồn tại. Hãy sao lưu và cập nhật có chọn lọc để giữ các tab/nội dung bổ sung.')
    profiles = json.loads(DATA.read_text())
    CACHE.mkdir(parents=True, exist_ok=True)
    urls = sorted({p['suppliedImageUrl'] for p in profiles if p.get('suppliedImageUrl')})
    # Reuse the 548 already downloaded covers.
    for row, p in enumerate(unique_games(profiles), start=2):
        old = ROOT / f'game-steam/exports/embedded-covers/{row:04d}.jpg'
        dest = path_for(p['suppliedImageUrl'])
        if old.exists() and not dest.exists():
            with Image.open(old) as image:
                image.thumbnail((160,160))
                image.save(dest, 'JPEG', quality=85)
    images, failures = {}, []
    with ThreadPoolExecutor(max_workers=16) as pool:
        pending = {pool.submit(fetch, url): url for url in urls}
        for future in as_completed(pending):
            url = pending[future]
            try:
                images[url], _ = future.result()
            except Exception as error:
                failures.append({'url': url, 'error': str(error)})
            done = len(images) + len(failures)
            if done % 250 == 0 or done == len(urls):
                print(f'Images {done}/{len(urls)}, downloaded {len(images)}, failures {len(failures)}', flush=True)
    REPORT.write_text(json.dumps({'total': len(urls), 'downloaded': len(images), 'failures': failures}, ensure_ascii=False, indent=2))
    genres = sorted({g for p in profiles for g in p.get('genres', [])})
    groups = [('Tất cả game', profiles), ('2 người - tay cầm', unique_games(profiles)),
              ('Hỗ trợ tay cầm', [p for p in profiles if nested(p,'steamMetadata','controllerSupport','level') in ('full','partial')]),
              ('Local Co-op', [p for p in profiles if nested(p,'steamMetadata','playerModes','localCoop') is True]),
              ('Online Co-op', [p for p in profiles if nested(p,'steamMetadata','playerModes','onlineCoop') is True])]
    groups += [(('Trực tuyến nhiều người' if g == 'Massively Multiplayer' else GENRE_VI.get(g,g)[:31]), [p for p in profiles if g in p.get('genres',[])]) for g in genres]
    groups += [('Chưa rõ thể loại', [p for p in profiles if not p.get('genres')]), ('Ảnh cần bổ sung', [p for p in profiles if p.get('suppliedImageUrl') not in images])]
    temp = OUTPUT.with_suffix('.building.xlsx')
    wb = xlsxwriter.Workbook(temp, {'strings_to_formulas': False, 'strings_to_urls': False})
    wb.set_properties({'title': 'Thư viện game tổng', 'comments': 'File chính để tiếp tục cập nhật; dữ liệu nguồn game-profiles.json.'})
    header = wb.add_format({'bold': True, 'bg_color': '#1F4E78', 'font_color': 'white', 'text_wrap': True, 'valign': 'vcenter'})
    body = wb.add_format({'text_wrap': True, 'valign': 'vcenter'})
    link = wb.add_format({'font_color': 'blue', 'underline': 1, 'valign': 'vcenter'})
    intro = wb.add_worksheet('Mục lục')
    intro.set_column('A:A', 38); intro.set_column('B:B', 110)
    notes = [
      ['THƯ VIỆN GAME TỔNG', 'File chính: Thu-vien-game-TONG.xlsx'],
      ['Tổng hồ sơ', len(profiles)], ['Ảnh nguồn tải được', len(images)], ['Ảnh chưa tải được', len(failures)],
      ['Cách phân loại', 'Một game có thể xuất hiện trong nhiều tab thể loại. Tab Tất cả game giữ đủ hồ sơ đầu vào, gồm cả tên/phiên bản trùng App ID.'],
      ['2 người - tay cầm', '548 tên duy nhất có hỗ trợ tay cầm full/partial và Local Co-op hoặc Shared/Split Screen. Steam API chưa xác nhận chính xác số tay cầm tối đa; cần kiểm tra từng game khi cần.'],
      ['Ảnh', 'Ảnh được nhúng trực tiếp, không dùng hàm IMAGE và xem được offline. Dòng thiếu ảnh được ghi rõ.'],
      ['Chưa rõ thể loại', 'Giữ riêng hồ sơ thiếu thể loại; không tự suy đoán thể loại.'],
      ['Cập nhật sau này', 'Tiếp tục dùng chính file này; giữ Mã hồ sơ để nối thông tin. Trước mỗi lần cập nhật hãy tạo bản sao lưu và giữ các tab/cột do người dùng bổ sung.'],
      ['File 548 game cũ', 'Giữ nguyên: 548-game-local-2-tay-cam-nhung-anh.xlsx'],
      ['Dữ liệu chi tiết', 'Tab Dữ liệu chi tiết giữ đầy đủ 30 cột nguồn, gồm trạng thái xác minh, dung lượng và nguồn kiểm chứng.'],
      ['Nguồn dữ liệu hiện tại', 'Steam Store đã tra cứu ngày 15/10/2026; đây không phải dữ liệu cập nhật trực tiếp.'],
      ['DANH SÁCH TAB', 'Số hồ sơ'],
    ]
    for row, values in enumerate(notes):
        intro.write_row(row,0,values,header if row in (0,12) else body)
        intro.set_row(row, 45 if row in (4,5,8) else 32)
    intro.freeze_panes(1,0)
    expected = {}
    for number, (name, records) in enumerate(groups):
        print(f'Building {name}: {len(records)}', flush=True)
        ws = wb.add_worksheet(name)
        expected[name] = len(records)
        r = len(notes) + number
        intro.write_url(r,0,"internal:'"+name.replace("'","''")+"'!A1",link,name)
        intro.write_number(r,1,len(records))
        ws.write_row(0,0,HEADERS,header); ws.set_row(0,36)
        ws.set_column(0,0,7); ws.set_column(1,1,16); ws.set_column(2,2,44)
        ws.set_column(3,3,38); ws.set_column(4,9,23); ws.set_column(10,11,48); ws.set_column(12,13,25)
        ws.freeze_panes(1,3); ws.autofilter(0,0,len(records),len(HEADERS)-1)
        ws.set_default_row(80)
        for r,p in enumerate(records,1):
            url = p.get('suppliedImageUrl','')
            app = nested(p,'candidateIdentifiers','steamAppId')
            store = nested(p,'steamMetadata','links','store') or (f'https://store.steampowered.com/app/{app}/' if app else '')
            controller = nested(p,'steamMetadata','controllerSupport','level')
            values = [r, '' if url in images else 'Chưa tải được ảnh', p.get('officialName') or p.get('suppliedName'),
                      '; '.join(GENRE_VI.get(g,g) for g in p.get('genres',[])),
                      {'full':'Đầy đủ','partial':'Một phần'}.get(controller,'Chưa xác nhận'),
                      *[flag(nested(p,'steamMetadata','playerModes',k)) for k in ('localCoop','sharedOrSplitScreen','onlineCoop')],
                      nested(p,'steamMetadata','storageRequirements','windows','minimum','display'),p.get('platforms',[]),
                      store,url,p.get('inputRecordId'),p.get('status')]
            ws.write_row(r,0,[clean(v) for v in values],body)
            if store: ws.write_url(r,10,store,link,'Mở Steam')
            if url: ws.write_url(r,11,url,link,'Ảnh nguồn')
            if url in images:
                path = images[url]
                with Image.open(path) as image:
                    width,height = image.size
                scale = min(100/height,100/width)
                ws.insert_image(r,1,path,{'x_scale':scale,'y_scale':scale,'x_offset':4,'y_offset':3,'object_position':1,'description':values[2]})
    detail = wb.add_worksheet('Dữ liệu chi tiết')
    detail.write_row(0,0,DETAIL_HEADERS,header)
    detail.set_column(0,29,26); detail.set_column(2,2,44)
    detail.freeze_panes(1,3); detail.autofilter(0,0,len(profiles),29)
    for r,p in enumerate(profiles,1):
        detail.write_row(r,0,[clean(v) for v in profile_row(r,p)],body)
    wb.close()
    with zipfile.ZipFile(temp) as z:
        assert z.testzip() is None
        print('Embedded media:',len([n for n in z.namelist() if n.startswith('xl/media/')]),flush=True)
    from openpyxl import load_workbook
    check = load_workbook(temp,read_only=True)
    for name,count in expected.items():
        assert check[name].max_row == count + 1, name
    assert check['Dữ liệu chi tiết'].max_row == len(profiles) + 1
    check.close()
    temp.replace(OUTPUT)
    manifest = {'workbook':str(OUTPUT),'profiles':len(profiles),'sheets':expected,'images_downloaded':len(images),'image_failures':len(failures),'source':str(DATA)}
    (CACHE.parent/'master-workbook-manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2))
    print(f'VALIDATED {OUTPUT}: {OUTPUT.stat().st_size:,} bytes',flush=True)

if __name__ == '__main__':
    main()
