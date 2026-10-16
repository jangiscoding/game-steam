#!/usr/bin/env python3
"""Download and embed Steam covers into the existing local-controller workbook."""
import io
import json
import time
import urllib.request
from concurrent.futures import ThreadPoolExecutor, as_completed
from pathlib import Path
from PIL import Image as PILImage
from openpyxl import load_workbook
from openpyxl.drawing.image import Image

ROOT = Path('/workspace')
SOURCE = ROOT / '548-game-local-2-tay-cam-co-hinh.xlsx'
OUTPUT = ROOT / '548-game-local-2-tay-cam-nhung-anh.xlsx'
CACHE = ROOT / 'game-steam/exports/embedded-covers'

def download(item):
    row, url = item
    path = CACHE / f'{row:04d}.jpg'
    if path.exists():
        with PILImage.open(path) as image:
            image.verify()
        return row, path
    for attempt in range(3):
        try:
            request = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0'})
            with urllib.request.urlopen(request, timeout=40) as response:
                data = response.read()
            with PILImage.open(io.BytesIO(data)) as image:
                image = image.convert('RGB')
                image.thumbnail((240, 240))
                image.save(path, 'JPEG', quality=88)
            return row, path
        except Exception:
            if attempt == 2:
                raise
            time.sleep(attempt + 1)

def main():
    CACHE.mkdir(parents=True, exist_ok=True)
    workbook = load_workbook(SOURCE)
    sheet = workbook['548 game local']
    items = [(r, sheet.cell(r, 10).value) for r in range(2, sheet.max_row + 1)]
    assert len(items) == 548 and all(url for _, url in items)
    images = {}
    failures = []
    with ThreadPoolExecutor(max_workers=12) as pool:
        pending = {pool.submit(download, item): item for item in items}
        for future in as_completed(pending):
            row, url = pending[future]
            try:
                _, path = future.result()
                images[row] = path
            except Exception as error:
                failures.append({'row': row, 'url': url, 'error': str(error)})
            done = len(images) + len(failures)
            if done % 50 == 0 or done == len(items):
                print(f'Downloaded {len(images)}/{len(items)}; failures {len(failures)}', flush=True)
    (CACHE.parent / 'image-download-report.json').write_text(json.dumps(failures, ensure_ascii=False, indent=2))
    if failures:
        raise RuntimeError(f'{len(failures)} image downloads failed; see report')
    for row, path in sorted(images.items()):
        sheet.cell(row, 2).value = None
        image = Image(path)
        scale = min(120 / image.height, 130 / image.width)
        image.width *= scale
        image.height *= scale
        sheet.add_image(image, f'B{row}')
        sheet.row_dimensions[row].height = 94
    for row in workbook['Tổng hợp'].iter_rows():
        if row[0].value == 'Lưu ý về ảnh':
            row[1].value = 'Đã nhúng trực tiếp đủ 548 ảnh vào file Excel. Ảnh hiển thị khi mở file, không cần hàm IMAGE hoặc kết nối Internet. Cột URL ảnh giữ đường dẫn nguồn Steam.'
    workbook.save(OUTPUT)
    check = load_workbook(OUTPUT)
    assert len(check['548 game local']._images) == 548
    assert check['548 game local'].max_row == 549
    assert not any(c.data_type == 'f' and 'IMAGE(' in str(c.value) for s in check for row in s for c in row)
    check.close()
    print(f'Validated 548 embedded images: {OUTPUT} ({OUTPUT.stat().st_size:,} bytes)', flush=True)

if __name__ == '__main__':
    main()
