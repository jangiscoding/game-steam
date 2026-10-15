#!/usr/bin/env python3
import argparse
import json
from collections import Counter
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


HEADERS = [
    "STT",
    "Input Record ID",
    "Tên đầu vào",
    "Steam App ID ứng viên",
    "Steam Store Link",
    "Tên chính thức",
    "Loại sản phẩm",
    "Edition",
    "Năm phát hành phiên bản",
    "Nhà phát triển",
    "Nhà phát hành",
    "Nền tảng",
    "Thể loại",
    "Dung lượng tối thiểu Windows",
    "Dung lượng khuyến nghị Windows",
    "Hỗ trợ controller",
    "Tay cầm Xbox",
    "Single-player",
    "Multiplayer",
    "Co-op",
    "Online Co-op",
    "Local Co-op",
    "Split-screen",
    "Số người chơi tối đa",
    "Trạng thái lấy dữ liệu Steam",
    "Lỗi lấy dữ liệu Steam",
    "Độ chắc chắn",
    "Trạng thái",
    "Ghi chú xác minh",
    "Nguồn kiểm chứng",
]


def parse_args():
    parser = argparse.ArgumentParser(description="Export game research profiles to Excel")
    parser.add_argument(
        "--input",
        default="data/game-profiles.json",
        help="Path to game-profiles.json",
    )
    parser.add_argument(
        "--output",
        default="steam-game-library.xlsx",
        help="Destination .xlsx path",
    )
    return parser.parse_args()


def clean(value):
    if value is None:
        return ""
    if isinstance(value, bool):
        return "Có" if value else "Không"
    if isinstance(value, (list, tuple)):
        value = "; ".join(str(item) for item in value if item not in (None, ""))
    return ILLEGAL_CHARACTERS_RE.sub("", str(value))


def nested(data, *keys):
    current = data
    for key in keys:
        if not isinstance(current, dict):
            return None
        current = current.get(key)
    return current


def storage_display(profile, level):
    return nested(
        profile,
        "steamMetadata",
        "storageRequirements",
        "windows",
        level,
        "display",
    )


def player_mode(profile, key):
    return nested(profile, "steamMetadata", "playerModes", key)


def source_links(profile):
    return [source.get("url") for source in profile.get("sources", []) if source.get("url")]


def xbox_value(profile):
    value = nested(profile, "steamMetadata", "controllerSupport", "xboxCompatible")
    if value is None:
        return "Chưa xác định"
    return "Có" if value else "Không"


def profile_row(index, profile):
    app_id = nested(profile, "candidateIdentifiers", "steamAppId")
    store_link = nested(profile, "steamMetadata", "links", "store")
    if not store_link and app_id:
        store_link = f"https://store.steampowered.com/app/{app_id}/"

    return [
        index,
        profile.get("inputRecordId"),
        profile.get("suppliedName"),
        app_id,
        store_link,
        profile.get("officialName"),
        profile.get("productType"),
        profile.get("editionName"),
        profile.get("identifiedVersionReleaseYear"),
        profile.get("developers", []),
        profile.get("publishers", []),
        profile.get("platforms", []),
        profile.get("genres", []),
        storage_display(profile, "minimum"),
        storage_display(profile, "recommended"),
        nested(profile, "steamMetadata", "controllerSupport", "level"),
        xbox_value(profile),
        player_mode(profile, "singlePlayer"),
        player_mode(profile, "multiplayer"),
        player_mode(profile, "coop"),
        player_mode(profile, "onlineCoop"),
        player_mode(profile, "localCoop"),
        player_mode(profile, "sharedOrSplitScreen"),
        player_mode(profile, "maximumPlayers"),
        nested(profile, "research", "steam", "status"),
        nested(profile, "research", "steam", "error"),
        profile.get("confidence"),
        profile.get("status"),
        profile.get("verificationNotes", []),
        source_links(profile),
    ]


def build_games_sheet(workbook, profiles):
    sheet = workbook.active
    sheet.title = "Games"
    sheet.append(HEADERS)

    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    status_fills = {
        "identified": PatternFill("solid", fgColor="C6EFCE"),
        "needs_verification": PatternFill("solid", fgColor="FFF2CC"),
        "unidentified": PatternFill("solid", fgColor="F4CCCC"),
    }

    for index, profile in enumerate(profiles, start=1):
        sheet.append([clean(value) for value in profile_row(index, profile)])
        row = sheet.max_row
        link_cell = sheet.cell(row=row, column=5)
        if link_cell.value:
            link_cell.hyperlink = link_cell.value
            link_cell.style = "Hyperlink"
        status = profile.get("status")
        if status in status_fills:
            sheet.cell(row=row, column=28).fill = status_fills[status]

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{sheet.max_row}"
    sheet.row_dimensions[1].height = 42

    widths = {
        1: 8,
        2: 22,
        3: 42,
        4: 22,
        5: 48,
        6: 42,
        7: 18,
        8: 22,
        9: 20,
        10: 32,
        11: 32,
        12: 28,
        13: 32,
        14: 24,
        15: 28,
        16: 22,
        17: 20,
        18: 16,
        19: 16,
        20: 14,
        21: 16,
        22: 16,
        23: 16,
        24: 22,
        25: 26,
        26: 48,
        27: 18,
        28: 22,
        29: 60,
        30: 50,
    }
    for column, width in widths.items():
        sheet.column_dimensions[get_column_letter(column)].width = width

    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def build_summary_sheet(workbook, profiles):
    sheet = workbook.create_sheet("Summary")
    status_counts = Counter(profile.get("status", "unknown") for profile in profiles)
    confidence_counts = Counter(profile.get("confidence", "unknown") for profile in profiles)
    researched = sum(
        nested(profile, "research", "steam", "status") == "success"
        for profile in profiles
    )
    steam_status_counts = Counter(
        nested(profile, "research", "steam", "status") or "pending"
        for profile in profiles
    )

    rows = [
        ("Chỉ số", "Giá trị"),
        ("Tổng hồ sơ", len(profiles)),
        ("Đã nghiên cứu Steam", researched),
        ("Đang chờ nghiên cứu Steam", steam_status_counts["pending"]),
        ("Chờ thử lại do rate limit", steam_status_counts["retry_pending"]),
        ("Steam không có dữ liệu công khai", steam_status_counts["not_found"]),
        ("Lỗi lấy dữ liệu Steam", steam_status_counts["failed"]),
        ("Đã xác định", status_counts["identified"]),
        ("Cần xác minh", status_counts["needs_verification"]),
        ("Chưa xác định", status_counts["unidentified"]),
        ("Độ chắc chắn cao", confidence_counts["high"]),
        ("Độ chắc chắn trung bình", confidence_counts["medium"]),
        ("Độ chắc chắn thấp", confidence_counts["low"]),
        ("", ""),
        (
            "Lưu ý",
            "Steam App ID ban đầu được trích từ URL ảnh trong dữ liệu đầu vào. "
            "Các dòng success đã được đối chiếu với Steam Store; retry_pending sẽ được thử lại sau khi hết rate limit.",
        ),
        (
            "Dung lượng",
            "Là dung lượng trống tối thiểu/khuyến nghị do Steam công bố, không phải kích thước tải xuống chính xác.",
        ),
        (
            "Controller/Số người chơi",
            "Không suy đoán Xbox compatibility hoặc số người chơi tối đa khi Steam không cung cấp bằng chứng rõ ràng.",
        ),
    ]
    for row in rows:
        sheet.append(row)
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    sheet.column_dimensions["A"].width = 32
    sheet.column_dimensions["B"].width = 110
    for row in sheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def build_dictionary_sheet(workbook):
    sheet = workbook.create_sheet("Data Dictionary")
    sheet.append(["Cột", "Ý nghĩa"])
    definitions = {
        "Steam App ID ứng viên": "ID chưa được coi là xác minh nếu chưa đối chiếu nguồn Steam chính thức.",
        "Tên chính thức": "Tên trả về từ nguồn chính thức sau khi nghiên cứu.",
        "Loại sản phẩm": "Game, DLC, remake, remaster, edition hoặc loại khác có bằng chứng.",
        "Dung lượng": "Required free space trong yêu cầu hệ thống của Steam.",
        "Hỗ trợ controller": "full, partial hoặc unknown theo dữ liệu Steam.",
        "Tay cầm Xbox": "Chỉ Có/Không khi nguồn ghi rõ; nếu không là Chưa xác định.",
        "Số người chơi tối đa": "Để trống nếu Steam app-details không cung cấp con số đáng tin cậy.",
        "Trạng thái lấy dữ liệu Steam": "success, pending, retry_pending, not_found hoặc failed.",
        "Độ chắc chắn": "high, medium hoặc low.",
        "Trạng thái": "identified, needs_verification hoặc unidentified.",
        "Nguồn kiểm chứng": "Các URL đã thực sự dùng để xác minh metadata.",
    }
    for column, meaning in definitions.items():
        sheet.append([column, meaning])
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    sheet.column_dimensions["A"].width = 35
    sheet.column_dimensions["B"].width = 100
    for row in sheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def main():
    args = parse_args()
    input_path = Path(args.input).resolve()
    output_path = Path(args.output).resolve()

    with input_path.open("r", encoding="utf-8") as handle:
        profiles = json.load(handle)

    workbook = Workbook()
    build_games_sheet(workbook, profiles)
    build_summary_sheet(workbook, profiles)
    build_dictionary_sheet(workbook)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)

    validated = load_workbook(output_path, read_only=True, data_only=False)
    games_sheet = validated["Games"]
    if games_sheet.max_row != len(profiles) + 1:
        raise RuntimeError("Export row count validation failed")
    validated.close()
    print(f"Exported {len(profiles)} profiles to {output_path}")


if __name__ == "__main__":
    main()
